from datetime import datetime
from io import BytesIO
import unittest
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from backend.settlement_export import build_settlement_workbook, settlement_export_filename


class SettlementExportTests(unittest.TestCase):
    def setUp(self):
        self.transactions = [
            {
                "id": "tx-20260916-001",
                "createdAt": 1_789_491_845_000,
                "payee": {
                    "username": "receiver",
                    "displayName": "收\x00款人甲",
                },
                "items": [
                    {
                        "category": "aws",
                        "consumptionAmount": "1234.5678",
                        "ratePercent": "45.5",
                        "settlementAmount": "561.7283",
                    },
                    {
                        "category": "openai",
                        "consumptionAmount": "20",
                        "ratePercent": "10",
                        "settlementAmount": "2",
                    },
                ],
                "totalSettlementAmount": "563.7283",
            },
            {
                "id": "=unsafe-id",
                "createdAt": 1_789_491_900_000,
                "payee": None,
                "items": [
                    {
                        "category": "anthropic",
                        "consumptionAmount": "5",
                        "ratePercent": "0",
                        "settlementAmount": "0",
                    }
                ],
                "totalSettlementAmount": "0",
            },
            {
                "id": "tx-20260916-003",
                "createdAt": 1_789_491_960_000,
                "payee": {"username": "receiver-2"},
                "items": [
                    {
                        "category": "vertexai",
                        "consumptionAmount": "144",
                        "ratePercent": "45",
                        "settlementAmount": "64.8",
                    }
                ],
                "totalSettlementAmount": "64.8",
            },
        ]

    def test_workbook_contains_summary_and_detail_records(self):
        content = build_settlement_workbook(
            self.transactions,
            username="xiaozhushou",
            user_id=34,
            language="zh",
        )
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["结算记录"])

        summary = workbook["结算记录"]
        self.assertEqual(
            [summary.cell(5, column).value for column in range(1, 9)],
            [
                "交易编号",
                "收款人",
                "交易结算总额",
                "结算时间",
                "渠道分类",
                "小号额度",
                "结算汇率",
                "明细结算金额",
            ],
        )
        self.assertEqual(summary["A6"].value, "tx-20260916-001")
        self.assertEqual(summary["B6"].value, "收款人甲 (receiver)")
        self.assertAlmostEqual(summary["C6"].value, 563.7283)
        self.assertIsInstance(summary["D6"].value, datetime)
        self.assertEqual(summary["E6"].value, "AWS Bedrock")
        self.assertAlmostEqual(summary["F6"].value, 1234.5678)
        self.assertAlmostEqual(summary["G6"].value, 0.455)
        self.assertAlmostEqual(summary["H6"].value, 561.7283)

        self.assertEqual(summary["A7"].value, "tx-20260916-001")
        self.assertEqual(summary["B7"].value, "收款人甲 (receiver)")
        self.assertAlmostEqual(summary["C7"].value, 563.7283)
        self.assertEqual(summary["E7"].value, "OpenAI")
        self.assertAlmostEqual(summary["H6"].value + summary["H7"].value, summary["C6"].value)

        self.assertEqual(summary["A8"].value, "'=unsafe-id")
        self.assertEqual(summary["A8"].data_type, "s")
        self.assertEqual(summary["B8"].value, "未记录")
        self.assertEqual(summary["A9"].value, "tx-20260916-003")
        self.assertEqual(summary["E9"].value, "Vertex AI Gemini")
        self.assertEqual(summary.max_row, 9)
        self.assertEqual(summary.freeze_panes, "A6")
        self.assertEqual(summary.auto_filter.ref, "A5:H9")
        self.assertTrue(all(cell.hyperlink is None for row in summary.iter_rows() for cell in row))

        for column in range(1, 9):
            self.assertEqual(
                summary.cell(6, column).fill.fgColor.rgb,
                summary.cell(7, column).fill.fgColor.rgb,
            )
            self.assertNotEqual(
                summary.cell(7, column).fill.fgColor.rgb,
                summary.cell(8, column).fill.fgColor.rgb,
            )
            self.assertNotEqual(
                summary.cell(8, column).fill.fgColor.rgb,
                summary.cell(9, column).fill.fgColor.rgb,
            )
        self.assertEqual(
            len(
                {
                    summary["A6"].fill.fgColor.rgb,
                    summary["A8"].fill.fgColor.rgb,
                    summary["A9"].fill.fgColor.rgb,
                }
            ),
            3,
        )
        self.assertEqual(summary["A6"].border.top.style, "medium")
        self.assertEqual(summary["A6"].border.bottom.style, "thin")
        self.assertEqual(summary["A7"].border.bottom.style, "medium")
        self.assertEqual(summary["A8"].border.top.style, "medium")
        self.assertEqual(summary["A8"].border.bottom.style, "medium")
        self.assertEqual(summary["A9"].border.top.style, "medium")
        self.assertEqual(summary["A9"].border.bottom.style, "medium")
        self.assertNotEqual(summary["A6"].border.top.color.rgb, summary["A8"].border.top.color.rgb)
        self.assertNotEqual(summary["A8"].border.top.color.rgb, summary["A9"].border.top.color.rgb)
        self.assertEqual(summary["A7"].border.left.style, "medium")
        self.assertEqual(summary["H7"].border.right.style, "medium")
        self.assertEqual(len(summary.merged_cells.ranges), 0)

    def test_english_workbook_and_filename(self):
        content = build_settlement_workbook(
            self.transactions[:1],
            username="receiver",
            user_id=35,
            language="en",
        )
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, ["Settlements"])
        self.assertEqual(workbook["Settlements"]["E6"].value, "AWS Bedrock")
        self.assertEqual(workbook["Settlements"]["H5"].value, "Detail Amount")
        filename = settlement_export_filename(
            "receiver",
            datetime(2026, 9, 16, 12, 34, 56, tzinfo=ZoneInfo("Asia/Shanghai")),
        )
        self.assertEqual(filename, "settlements-receiver-20260916-123456.xlsx")

    def test_order_palette_keeps_adjacent_transactions_distinct(self):
        transactions = []
        for index in range(9):
            transactions.append(
                {
                    "id": f"tx-{index + 1}",
                    "createdAt": 1_789_491_845_000 + index * 1_000,
                    "payee": {"username": "receiver"},
                    "items": [
                        {
                            "category": "aws",
                            "consumptionAmount": "1",
                            "ratePercent": "45",
                            "settlementAmount": "0.45",
                        }
                    ],
                    "totalSettlementAmount": "0.45",
                }
            )

        content = build_settlement_workbook(
            transactions,
            username="receiver",
            user_id=35,
            language="zh",
        )
        sheet = load_workbook(BytesIO(content))["结算记录"]
        fills = [sheet.cell(6 + index, 1).fill.fgColor.rgb for index in range(9)]

        self.assertEqual(len(set(fills[:8])), 8)
        self.assertEqual(fills[8], fills[0])
        for index in range(8):
            self.assertNotEqual(fills[index], fills[index + 1])
        for row in range(6, 15):
            self.assertEqual(sheet.cell(row, 1).fill.fill_type, "solid")
            self.assertEqual(sheet.cell(row, 1).border.top.style, "medium")
            self.assertEqual(sheet.cell(row, 1).border.bottom.style, "medium")
            self.assertEqual(
                sheet.cell(row, 1).border.top.color.rgb,
                sheet.cell(row, 1).border.bottom.color.rgb,
            )

    def test_missing_detail_data_is_rejected(self):
        invalid_cases = [
            [{**self.transactions[0], "items": []}],
            [{**self.transactions[0], "items": [None]}],
            [{**self.transactions[0], "items": [{"category": "aws", "consumptionAmount": None, "ratePercent": "10", "settlementAmount": "1"}]}],
        ]
        for transactions in invalid_cases:
            with self.subTest(transactions=transactions):
                with self.assertRaises(ValueError):
                    build_settlement_workbook(
                        transactions,
                        username="receiver",
                        user_id=35,
                        language="zh",
                    )


if __name__ == "__main__":
    unittest.main()
