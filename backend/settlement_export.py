from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


GYS_TIMEZONE = ZoneInfo("Asia/Shanghai")

CATEGORY_LABELS_ZH = {
    "aws": "AWS Bedrock",
    "aws_a": "Claude on AWS",
    "anthropic": "Anthropic 官方",
    "anthropic_small": "Anthropic 小额",
    "anthropic_test": "Anthropic 测试小额",
    "anthropic_ent": "Anthropic 参数覆盖",
    "openai": "OpenAI",
    "azure": "Azure OpenAI",
    "azure_claude": "Azure Claude",
    "ai_studio": "Google AI Studio",
    "vertexai": "Vertex AI Gemini",
    "vertexai_claude": "Vertex AI Claude",
    "openrouter": "OpenRouter",
    "opencode": "OpenCode Claude",
}

CATEGORY_LABELS_EN = {
    "aws": "AWS Bedrock",
    "aws_a": "Claude on AWS",
    "anthropic": "Anthropic Official",
    "anthropic_small": "Anthropic Small",
    "anthropic_test": "Anthropic Test Small",
    "anthropic_ent": "Anthropic Parameter Override",
    "openai": "OpenAI",
    "azure": "Azure OpenAI",
    "azure_claude": "Azure Claude",
    "ai_studio": "Google AI Studio",
    "vertexai": "Vertex AI Gemini",
    "vertexai_claude": "Vertex AI Claude",
    "openrouter": "OpenRouter",
    "opencode": "OpenCode Claude",
}

HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="10233F")
BODY_FONT = Font(name="Arial", size=10, color="1F2937")
CONTEXT_FONT = Font(name="Arial", size=9, italic=True, color="64748B")
ORDER_PALETTE = (
    ("E8F1FF", "7FA6D8"),
    ("EAF7EF", "7FB38E"),
    ("FFF1D8", "D4A455"),
    ("F2EAFF", "A58ACB"),
    ("FFE8EC", "D78C9A"),
    ("E6F7F6", "72B8B3"),
    ("F4F0E8", "B8A27D"),
    ("EBEFF5", "8C9BB0"),
)
ORDER_ROW_SIDE = Side(style="thin", color="D9E2EF")
MONEY_FORMAT = '$#,##0.0000'
PERCENT_FORMAT = "0.00%"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"


def safe_excel_text(value: Any) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def decimal_value(value: Any) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError("Missing settlement amount")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise ValueError("Invalid settlement amount") from error
    if not result.is_finite():
        raise ValueError("Invalid settlement amount")
    return result


def beijing_datetime(timestamp_ms: Any) -> datetime:
    try:
        timestamp = int(timestamp_ms)
    except (TypeError, ValueError) as error:
        raise ValueError("Invalid settlement timestamp") from error
    return (
        datetime.fromtimestamp(timestamp / 1000, tz=timezone.utc)
        .astimezone(GYS_TIMEZONE)
        .replace(tzinfo=None)
    )


def person_label(person: Any, fallback: str = "") -> str:
    if not isinstance(person, dict):
        return safe_excel_text(fallback)
    username = str(person.get("username") or "").strip()
    display_name = str(person.get("displayName") or "").strip()
    if display_name and username and display_name != username:
        return safe_excel_text(f"{display_name} ({username})")
    return safe_excel_text(display_name or username or fallback)


def style_sheet(sheet: Any, widths: tuple[float, ...]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A6"
    sheet.row_dimensions[2].height = 22
    sheet.row_dimensions[5].height = 23
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def style_header(sheet: Any, row: int, columns: int) -> None:
    sheet.row_dimensions[row].height = 23
    for cell in sheet[row][:columns]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def style_body(
    sheet: Any,
    transaction_rows: list[tuple[int, int]],
    columns: int,
) -> None:
    for transaction_index, (start_row, end_row) in enumerate(transaction_rows):
        fill_color, outline_color = ORDER_PALETTE[transaction_index % len(ORDER_PALETTE)]
        fill = PatternFill("solid", fgColor=fill_color)
        outline_side = Side(style="medium", color=outline_color)
        for row in range(start_row, end_row + 1):
            is_first_row = row == start_row
            is_last_row = row == end_row
            for column, cell in enumerate(sheet[row][:columns], start=1):
                cell.fill = fill
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="center")
                cell.border = Border(
                    left=outline_side if column == 1 else Side(),
                    right=outline_side if column == columns else Side(),
                    top=outline_side if is_first_row else Side(),
                    bottom=outline_side if is_last_row else ORDER_ROW_SIDE,
                )
            sheet.row_dimensions[row].height = 23 if is_first_row else 21

        for cell in sheet[start_row][:4]:
            cell.font = Font(name="Arial", size=10, bold=True, color="1F2937")


def build_settlement_workbook(
    transactions: list[dict[str, Any]],
    *,
    username: str,
    user_id: int,
    language: str = "zh",
) -> bytes:
    english = language == "en"
    labels = CATEGORY_LABELS_EN if english else CATEGORY_LABELS_ZH
    sheet_name = "Settlements" if english else "结算记录"
    summary_title = "Settlement records" if english else "结算记录"
    missing_payee = "Not recorded" if english else "未记录"
    context = (
        f"Account: {safe_excel_text(username)}  ·  User ID: {user_id}"
        if english
        else f"账户：{safe_excel_text(username)}  ·  用户 ID：{user_id}"
    )

    workbook = Workbook()
    workbook.properties.creator = "PushKey System"
    workbook.properties.title = summary_title
    sheet = workbook.active
    sheet.title = sheet_name

    sheet["A2"] = summary_title
    sheet["A2"].font = TITLE_FONT
    sheet["A3"] = context
    sheet["A3"].font = CONTEXT_FONT
    headers = (
        [
            "Transaction ID",
            "Payee",
            "Transaction Total",
            "Settlement Time",
            "Channel Category",
            "Sub-account Quota",
            "Settlement Rate",
            "Detail Amount",
        ]
        if english
        else [
            "交易编号",
            "收款人",
            "交易结算总额",
            "结算时间",
            "渠道分类",
            "小号额度",
            "结算汇率",
            "明细结算金额",
        ]
    )
    for column, value in enumerate(headers, start=1):
        sheet.cell(5, column, value)

    transaction_rows: list[tuple[int, int]] = []
    for transaction in transactions:
        transaction_id = safe_excel_text(transaction.get("id"))
        items = transaction.get("items") if isinstance(transaction.get("items"), list) else []
        if not items:
            raise ValueError("Settlement transaction has no detail records")
        payee = person_label(transaction.get("payee"), missing_payee)
        total_amount = decimal_value(transaction.get("totalSettlementAmount"))
        created_at = beijing_datetime(transaction.get("createdAt"))
        transaction_start_row = sheet.max_row + 1
        for item in items:
            if not isinstance(item, dict):
                raise ValueError("Invalid settlement detail record")
            category = str(item.get("category") or "")
            if not category:
                raise ValueError("Missing settlement category")
            sheet.append(
                [
                    transaction_id,
                    payee,
                    total_amount,
                    created_at,
                    labels.get(category, safe_excel_text(category)),
                    decimal_value(item.get("consumptionAmount")),
                    decimal_value(item.get("ratePercent")) / Decimal(100),
                    decimal_value(item.get("settlementAmount")),
                ]
            )
        transaction_rows.append((transaction_start_row, sheet.max_row))

    style_sheet(sheet, (40, 28, 22, 22, 28, 22, 20, 22))
    style_header(sheet, 5, 8)
    if sheet.max_row >= 6:
        style_body(sheet, transaction_rows, 8)
        sheet.auto_filter.ref = f"A5:H{sheet.max_row}"
        for row in sheet.iter_rows(min_row=6, max_row=sheet.max_row):
            row[2].number_format = MONEY_FORMAT
            row[3].number_format = DATETIME_FORMAT
            row[5].number_format = MONEY_FORMAT
            row[6].number_format = PERCENT_FORMAT
            row[7].number_format = MONEY_FORMAT

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def settlement_export_filename(username: str, now: datetime | None = None) -> str:
    current = (now or datetime.now(GYS_TIMEZONE)).astimezone(GYS_TIMEZONE)
    safe_username = "".join(character if character.isalnum() or character in "_.-" else "-" for character in username)
    return f"settlements-{safe_username or 'account'}-{current.strftime('%Y%m%d-%H%M%S')}.xlsx"
